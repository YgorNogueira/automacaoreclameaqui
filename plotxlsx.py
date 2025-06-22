import pandas as pd

# 0) (Opcional) verifique se openpyxl está instalado:
#    pip install openpyxl

# 1) Ler o CSV original
df = pd.read_csv('reclamacoes.csv')

# 2) Selecionar as colunas originais e renomeá-las
df_out = df[['titulo', 'descricao', 'resolvida', 'tempo_envio']].copy()
df_out.rename(columns={
    'titulo': 'título',
    'descricao': 'descrição',
    'tempo_envio': 'tempo de envio'
}, inplace=True)

# 3) Criar o arquivo .xlsx usando openpyxl
output_file = 'reclamacoes.xlsx'
with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
    # 3.1) Gravar o DataFrame na sheet "RECLAMAÇÕES"
    df_out.to_excel(writer, sheet_name='RECLAMAÇÕES', index=False)
    
    # 3.2) Formatar via openpyxl (obtendo workbook/worksheet)
    workbook  = writer.book
    worksheet = writer.sheets['RECLAMAÇÕES']
    
    # 4) Formatar cabeçalho: negrito, fundo e borda
    header_fmt = {
        'font': {'bold': True},
        'fill': {'fill_type': 'solid', 'start_color': 'D7E4BC'},
        'border': {
            'top':    {'style': 'thin'},
            'bottom': {'style': 'thin'},
            'left':   {'style': 'thin'},
            'right':  {'style': 'thin'},
        }
    }
    from openpyxl.styles import Font, PatternFill, Border, Side
    font = Font(bold=True)
    fill = PatternFill(fill_type='solid', start_color='D7E4BC')
    side = Side(style='thin')
    border = Border(top=side, bottom=side, left=side, right=side)
    
    for col_idx, _ in enumerate(df_out.columns, 1):
        cell = worksheet.cell(row=1, column=col_idx)
        cell.font = font
        cell.fill = fill
        cell.border = border
    
    # 5) Ajustar largura de cada coluna ao conteúdo
    for i, col in enumerate(df_out.columns, 1):
        max_len = max(
            df_out[col].astype(str).map(len).max(),
            len(col)
        ) + 2
        worksheet.column_dimensions[worksheet.cell(row=1, column=i).column_letter].width = max_len

# 6) Mostrar o DataFrame no console ao final
print(df_out)
